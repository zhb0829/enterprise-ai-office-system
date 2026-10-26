from io import BytesIO

import pytest
from fastapi import HTTPException
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.config import Settings
from app.db import Base
from app.models import MaterialChunk, PolicyClause, ReferenceMaterial
from app.services import materials, policy
from app.agent.governance import DataLevel, Principal, data_level_var, model_messages, principal_var
from app.agent.knowledge import KnowledgeStore
from app.agent.platform_models import KnowledgeRelease


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setattr(materials.settings, "embedding_api_key", "")
    monkeypatch.setattr(materials.settings, "chunk_semantic_enabled", True)
    monkeypatch.setattr(materials.settings, "chunk_semantic_trigger_size", 1200)
    monkeypatch.setattr(materials.settings, "chunk_min_size", 320)
    monkeypatch.setattr(materials.settings, "chunk_semantic_similarity_threshold", 0.42)
    token = data_level_var.set(DataLevel.public)
    principal_token = principal_var.set(Principal(
        user_id="tester", tenant_id="default", clearance=DataLevel.restricted,
        public_only=False,
    ))
    yield
    principal_var.reset(principal_token)
    data_level_var.reset(token)


@pytest.fixture
def database():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield session
    engine.dispose()


def topic_text():
    return ("研发数据" * 59 + "。") * 3 + ("预算审计" * 59 + "。") * 3


def topic_vectors(texts):
    return [[1.0, 0.0] if "研发" in text else [0.0, 1.0] for text in texts]


def test_short_text_does_not_call_semantic_embedder():
    def unexpected(_):
        pytest.fail("short documents must not invoke semantic embeddings")

    chunks = materials.split_text_chunks("完整的短段落。\n\n第二段。", embedder=unexpected)
    assert len(chunks) == 1
    assert chunks[0].split_method == "structure"


def test_topic_change_splits_once_and_retains_exact_source():
    text = topic_text()
    chunks = materials.split_text_chunks(text, embedder=topic_vectors)
    assert len(chunks) == 2
    assert all(chunk.split_method == "semantic" for chunk in chunks)
    assert all(chunk.parent_index == 0 and chunk.parent_text == text for chunk in chunks)
    assert "".join(chunk.text for chunk in chunks) == text
    assert all(text[chunk.start:chunk.end] == chunk.text for chunk in chunks)
    assert all(320 <= len(chunk.text) <= 900 for chunk in chunks)


def test_semantic_tail_is_not_a_tiny_fragment():
    text = ("研发数据" * 59 + "。") * 5 + "短结尾。"
    chunks = materials.split_text_chunks(
        text, chunk_size=2000,
        embedder=lambda texts: [[float(index % 2), float((index + 1) % 2)] for index in range(len(texts))],
    )
    assert all(len(chunk.text) >= 320 for chunk in chunks)


@pytest.mark.parametrize("invalid", [None, [], [[float("nan")]], [[0.0, 0.0]]])
def test_semantic_failure_and_invalid_vectors_use_recursive_fallback(invalid):
    text = topic_text()

    def failing(texts):
        if invalid is None:
            raise materials.EmbeddingError("offline")
        return invalid

    chunks = materials.split_text_chunks(text, embedder=failing)
    assert [chunk.text for chunk in chunks] == materials.split_text(text)
    assert all(chunk.split_method == "recursive" for chunk in chunks)
    assert all(text[chunk.start:chunk.end] == chunk.text for chunk in chunks)


def test_headings_do_not_merge_and_repeated_text_offsets_are_exact():
    paragraph = "同一段。" * 220
    text = f"# 第一章\n\n{paragraph}\n\n# 第二章\n\n{paragraph}"
    chunks = materials.split_text_chunks(text, semantic=False)
    assert not any("# 第一章" in chunk.text and "# 第二章" in chunk.text for chunk in chunks)
    assert all(len(chunk.text) <= 900 for chunk in chunks)
    assert all(text[chunk.start:chunk.end] == chunk.text for chunk in chunks)
    assert len({chunk.parent_index for chunk in chunks}) == 2


@pytest.mark.parametrize("text", [
    "| 项目 | 说明 |\n" + "| 研发 | 应当保留资料。 |\n" * 100,
    "- 申报条件：应当提交材料。\n" * 100,
    "```\n" + "return value;\n" * 100 + "```",
    "```python\n" + "return value;\n\n" * 100 + "```",
])
def test_tables_lists_and_code_are_not_semantically_split(text):
    chunks = materials.split_text_chunks(text, embedder=lambda _: pytest.fail("protected block"))
    assert len(chunks) == 1
    assert chunks[0].text == text.strip()


@pytest.mark.parametrize("size,overlap", [(0, 0), (5, 5), (5, -1)])
def test_invalid_sizes_are_rejected(size, overlap):
    with pytest.raises(ValueError):
        materials.split_text("text", chunk_size=size, overlap=overlap)


def test_docx_preserves_paragraph_and_table_order():
    from docx import Document

    doc = Document()
    doc.add_paragraph("前文")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    doc.add_paragraph("后文")
    buffer = BytesIO()
    doc.save(buffer)
    assert materials.parse_material("test.docx", buffer.getvalue()) == "前文\n\n| 字段 | 值 |\n\n后文"


def test_material_index_only_contains_children_and_retrieval_deduplicates_parents(database):
    text = topic_text()
    material = ReferenceMaterial(filename="topics.txt", text_content=text)
    database.add(material)
    database.flush()
    count = materials.build_material_chunks(database, material)
    database.commit()
    rows = database.query(MaterialChunk).all()
    assert count == len(rows) > 1
    assert all(row.meta["parent_text"] == text for row in rows)
    assert all(row.meta["chunking_version"] == "hybrid-v1" for row in rows)
    hits = materials.search_material_chunks(database, "研发数据")
    assert len(hits) == 1
    assert hits[0]["meta"]["parent_text"] == text


def test_semantic_and_index_embeddings_use_material_classification(database, monkeypatch):
    monkeypatch.setattr(materials.settings, "embedding_api_key", "configured")
    calls = []

    def embedding(texts):
        calls.append(data_level_var.get())
        return topic_vectors(texts)

    monkeypatch.setattr(materials, "embed_texts", embedding)
    material = ReferenceMaterial(filename="private.txt", classification="internal", text_content=topic_text())
    database.add(material)
    database.flush()
    materials.build_material_chunks(database, material)
    assert len(calls) == 2
    assert calls == [DataLevel.internal, DataLevel.internal]


def test_restricted_data_is_blocked_before_any_semantic_request(database, monkeypatch):
    monkeypatch.setattr(materials.settings, "embedding_api_key", "configured")
    calls = []

    def embedding(texts):
        model_messages([{"content": text} for text in texts])
        calls.append(texts)
        return topic_vectors(texts)

    monkeypatch.setattr(materials, "embed_texts", embedding)
    material = ReferenceMaterial(filename="restricted.txt", classification="restricted", text_content=topic_text())
    database.add(material)
    database.flush()
    with pytest.raises(HTTPException):
        materials.build_material_chunks(database, material)
    assert calls == []


def test_final_index_embedding_failure_is_not_hidden(database, monkeypatch):
    monkeypatch.setattr(materials.settings, "embedding_api_key", "configured")
    monkeypatch.setattr(
        materials, "embed_texts",
        lambda _: (_ for _ in ()).throw(materials.EmbeddingError("remote failed")),
    )
    material = ReferenceMaterial(filename="topics.txt", text_content=topic_text())
    database.add(material)
    database.flush()
    with pytest.raises(materials.EmbeddingError):
        materials.build_material_chunks(database, material)


def test_short_articles_remain_whole():
    text = "第一条 " + "应当保存研发数据。" * 110
    parts = policy._refine_policy_clauses(policy.split_policy_clauses(text))
    assert len(parts) == 1
    assert parts[0]["content"] == text
    assert parts[0]["meta"]["split_method"] == "article"


def test_policy_children_retrieve_and_cite_whole_parent(database):
    text = "第一条 " + topic_text()
    document = policy.create_policy_document(database, title="研发审计", content=text + "\n第二条 其他规定。")
    children = database.query(PolicyClause).filter_by(doc_id=document.id, article_no="第一条").all()
    assert len(children) > 1
    assert all(child.meta["parent_text"] == text for child in children)
    hits = policy.search_policy_clauses(database, "研发数据")
    first = [hit for hit in hits if hit["clause"].article_no == "第一条"]
    assert len(first) == 1
    assert policy.citation_for(first[0])["excerpt"] == text


def test_rechunk_is_opt_in_and_preserves_old_snapshot(database):
    text = "第一条 " + topic_text()
    document = policy.create_policy_document(database, title="研发审计", content=text)
    database.query(PolicyClause).filter_by(doc_id=document.id).delete(synchronize_session="fetch")
    database.add(PolicyClause(
        doc_id=document.id, article_no="第一条", content=text,
        embedding=materials.embed_text(text), meta=materials.embedding_info(materials.embed_text(text)),
    ))
    database.flush()
    policy.sync_policy_document(database, document)
    database.commit()
    old = KnowledgeStore(database, "public:policy").active()
    original_ids = [row.id for row in database.query(PolicyClause).all()]
    policy.reindex_policy_embeddings(database, document.id)
    assert [row.id for row in database.query(PolicyClause).all()] == original_ids
    result = policy.reindex_policy_embeddings(database, document.id, rechunk=True)
    assert result["clauses"] > 1
    assert database.get(KnowledgeRelease, old.id) is not None
    assert KnowledgeStore(database, "public:policy").active().id != old.id
    count = database.query(PolicyClause).count()
    release = KnowledgeStore(database, "public:policy").active().id
    policy.reindex_policy_embeddings(database, document.id, rechunk=True)
    assert database.query(PolicyClause).count() == count
    assert KnowledgeStore(database, "public:policy").active().id == release


def test_rechunk_embedding_failure_keeps_old_policy_rows_and_snapshot(database, monkeypatch):
    text = "第一条 " + topic_text()
    document = policy.create_policy_document(database, title="研发审计", content=text)
    old_rows = [(row.id, row.content) for row in database.query(PolicyClause).all()]
    old_release = KnowledgeStore(database, "public:policy").active().id
    monkeypatch.setattr(
        policy, "embed_texts",
        lambda _: (_ for _ in ()).throw(materials.EmbeddingError("remote failed")),
    )
    with pytest.raises(materials.EmbeddingError):
        policy.reindex_policy_embeddings(database, document.id, rechunk=True)
    database.rollback()
    assert [(row.id, row.content) for row in database.query(PolicyClause).all()] == old_rows
    assert KnowledgeStore(database, "public:policy").active().id == old_release


def test_raw_knowledge_uses_hybrid_metadata_without_repeating_embedding(database, monkeypatch):
    calls = []

    def embedding(texts):
        calls.append(data_level_var.get())
        return topic_vectors(texts)

    monkeypatch.setattr(materials, "embed_texts", embedding)
    monkeypatch.setattr(materials.settings, "embedding_api_key", "configured")
    store = KnowledgeStore(database, "tenant:default:materials")
    payload = [{"key": "raw", "text": topic_text(), "classification": "internal", "tenant_id": "default"}]
    row = store.stage(payload)
    store.publish(row.id)
    database.commit()
    assert calls == [DataLevel.internal, DataLevel.internal]
    chunks = store.snapshot()["entries"][0]["chunks"]
    assert len(chunks) == 2
    assert all(chunk["meta"]["parent_text"] == topic_text() for chunk in chunks)
    store.stage(payload)
    assert len(calls) == 2


def test_generic_knowledge_search_deduplicates_hybrid_parent(database, monkeypatch):
    monkeypatch.setattr(materials, "embed_texts", lambda texts: topic_vectors(texts))
    monkeypatch.setattr(materials.settings, "embedding_api_key", "configured")
    store = KnowledgeStore(database, "tenant:default:materials")
    row = store.stage([{
        "key": "raw-search", "text": topic_text(), "classification": "internal", "tenant_id": "default",
    }])
    store.publish(row.id)
    database.commit()
    from app.agent.knowledge import search_knowledge

    results = search_knowledge(database, "materials", "研发数据", limit=6)
    assert len(results) == 1
    assert results[0]["parentText"] == topic_text()


def test_chunking_configuration_rejects_bad_threshold():
    with pytest.raises(ValueError):
        Settings(_env_file=None, chunk_semantic_similarity_threshold=2)
