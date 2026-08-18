"""导出服务：草稿导出为 Markdown / DOCX。

PDF 导出依赖本机 pandoc + xelatex，中文字体可通过 PDF_CHINESE_FONT 配置。
"""
import logging
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from sqlalchemy.orm import Session

from ..config import settings
from ..models import Draft, ExportLog

logger = logging.getLogger(__name__)


class UnsupportedExportFormat(Exception):
    pass


def export_draft(db: Session, draft_id: str, fmt: str) -> tuple[str, str]:
    """导出草稿，返回 (文件路径, 文件名)。"""
    draft = db.get(Draft, draft_id)
    if not draft:
        raise FileNotFoundError(f"草稿不存在: {draft_id}")

    safe_title = re.sub(r'[\\/:*?"<>|]', "_", draft.title or draft.id)[:60] or "draft"
    filename = f"{safe_title}_{draft.id}.{fmt}"
    path = settings.export_path / filename

    if fmt == "md":
        path.write_text(draft.content, encoding="utf-8")
    elif fmt == "docx":
        _md_to_docx(draft.content, str(path))
    elif fmt == "pdf":
        _md_to_pdf(draft.content, path)
    else:
        raise UnsupportedExportFormat(f"不支持的导出格式: {fmt}")

    db.add(ExportLog(draft_id=draft_id, format=fmt, file_path=str(path)))
    db.commit()
    logger.info("导出成功: draft=%s fmt=%s path=%s", draft_id, fmt, path)
    return str(path), filename


def _clean_md(text: str) -> str:
    """去除 Markdown 强调符号，便于 DOCX 显示。"""
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("__", "").replace("~~", "").replace("*", "")


def _md_to_docx(md: str, path: str) -> None:
    """将 Markdown 草稿简单转换为 DOCX（标题/列表/段落）。"""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    # 中文默认字体（改善 DOCX 中文显示）
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "SimSun"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    in_code = False
    for line in md.splitlines():
        s = line.rstrip()
        if s.startswith("```"):
            in_code = not in_code
            continue
        if not s:
            doc.add_paragraph("")
            continue
        text = s if in_code else _clean_md(s)
        if s.startswith("### "):
            doc.add_heading(text[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(text[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(text[2:], level=1)
        elif s.startswith("- "):
            doc.add_paragraph(text[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s+", s):
            doc.add_paragraph(re.sub(r"^\d+\.\s+", "", text), style="List Number")
        elif s.startswith("> "):
            doc.add_paragraph(text[2:], style="Intense Quote")
        elif "|" in s and s.strip().startswith("|"):
            doc.add_paragraph(" ".join(part.strip() for part in text.strip("|").split("|") if part.strip()))
        else:
            doc.add_paragraph(text)
    doc.save(path)


def _md_to_pdf(md: str, path: Path) -> None:
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise UnsupportedExportFormat("PDF 导出需要安装 pandoc，并配置可用的 xelatex 与中文字体")

    with tempfile.TemporaryDirectory() as tmp:
        md_path = Path(tmp) / "draft.md"
        md_path.write_text(md, encoding="utf-8")
        cmd = [
            pandoc,
            str(md_path),
            "-o",
            str(path),
            "--pdf-engine=xelatex",
            "-V",
            f"CJKmainfont={settings.pdf_chinese_font}",
        ]
        try:
            subprocess.run(cmd, check=True, capture_output=True, text=True, encoding="utf-8")
        except FileNotFoundError as exc:
            raise UnsupportedExportFormat("PDF 导出需要安装 pandoc 与 xelatex") from exc
        except subprocess.CalledProcessError as exc:
            message = (exc.stderr or exc.stdout or str(exc)).strip()
            raise UnsupportedExportFormat(f"PDF 导出失败，请检查 pandoc/xelatex/中文字体配置: {message}") from exc
