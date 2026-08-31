"""公开会议资料解析：文本、文档、PPTX、图片、音频与链接正文。"""
from __future__ import annotations

import base64
import re
from pathlib import Path

from ..config import settings

MAX_AUDIO_BYTES = 500 * 1024 * 1024
MAX_FILE_BYTES = 100 * 1024 * 1024


def parse_transcript(text: str) -> dict:
    text = _normalize(text)
    segments = []
    speaker_pattern = re.compile(r"^\s*([^：:\n]{1,30})[：:]\s*(.+)$")
    for index, paragraph in enumerate(_paragraphs(text), start=1):
        match = speaker_pattern.match(paragraph)
        speaker = match.group(1).strip() if match else ""
        content = match.group(2).strip() if match else paragraph
        segments.append({
            "speaker": speaker,
            "text": content,
            "location": f"第{index}段",
        })
    return {"text": text, "segments": segments, "hasSpeakers": any(s["speaker"] for s in segments)}


def parse_material(material: dict) -> dict:
    if material.get("parseStatus") == "parsed" and (material.get("parseResult") or {}).get("text"):
        return material
    kind = str(material.get("materialType") or "")
    try:
        if kind == "link":
            result = parse_link(str(material.get("sourceUrl") or ""))
        else:
            path = Path(str(material.get("filePath") or ""))
            if not path.is_file():
                raise ValueError("资料文件不存在或 Python Worker 无权访问")
            size = int(material.get("sizeBytes") or path.stat().st_size)
            if kind == "audio":
                if size > MAX_AUDIO_BYTES:
                    raise ValueError("音频超过 500MB")
                result = parse_audio(path)
            else:
                if size > MAX_FILE_BYTES:
                    raise ValueError("文件超过 100MB")
                result = parse_file(path, str(material.get("mimeType") or ""), kind)
        material["parseStatus"] = "parsed"
        material["parseResult"] = result
    except Exception as exc:  # noqa: BLE001
        material["parseStatus"] = "failed"
        material["parseResult"] = {"error": str(exc)[:1000], "text": "", "segments": []}
    return material


def parse_link(url: str) -> dict:
    import trafilatura

    if not re.match(r"^https?://", url):
        raise ValueError("资料链接仅支持 http/https")
    downloaded = trafilatura.fetch_url(url)
    if not downloaded:
        raise ValueError("公开链接抓取失败")
    text = trafilatura.extract(downloaded, include_links=False, include_tables=True)
    if not text:
        raise ValueError("公开链接未提取到正文")
    return _paragraph_result(text, "原文段落")


def parse_file(path: Path, mime: str, kind: str) -> dict:
    ext = path.suffix.lower()
    if kind == "image" or mime.startswith("image/"):
        return parse_image(path, mime)
    if ext in {".txt", ".md"}:
        return parse_transcript(_decode(path.read_bytes()))
    if ext == ".pdf":
        return _parse_pdf(path)
    if ext == ".docx":
        return _parse_docx(path)
    if ext == ".pptx":
        return _parse_pptx(path)
    raise ValueError(f"不支持的会议资料格式：{ext or mime}")


def parse_image(path: Path, mime: str) -> dict:
    if not settings.meeting_vision_model:
        raise ValueError("未配置会议视觉模型，无法解析图片资料")
    vision_api_key = settings.meeting_vision_api_key or settings.llm_api_key
    if not vision_api_key:
        raise ValueError("未配置会议视觉模型 API Key，无法解析图片资料")
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    data_uri = f"data:{mime or 'image/png'};base64,{encoded}"
    from openai import OpenAI

    client = OpenAI(
        api_key=vision_api_key,
        base_url=settings.meeting_vision_base_url or settings.llm_base_url,
        timeout=settings.llm_timeout,
    )
    try:
        response = client.chat.completions.create(
            model=settings.meeting_vision_model,
            temperature=0,
            messages=[
                {"role": "system", "content": "提取会议/PPT截图中的标题、正文、数据、图表结论和来源说明。只输出可见内容，不推测。"},
                {"role": "user", "content": [
                    {"type": "text", "text": "请按原有层次转写图片中的会议信息。"},
                    {"type": "image_url", "image_url": {"url": data_uri}},
                ]},
            ],
        )
        text = response.choices[0].message.content or ""
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"视觉模型调用失败: {exc}") from exc
    if not text.strip():
        raise ValueError("视觉模型未返回可用文本")
    return _paragraph_result(text, "图片区域")


def parse_audio(path: Path) -> dict:
    if not settings.meeting_asr_api_key:
        raise ValueError("未配置会议 ASR 服务，无法转写音频")
    from openai import OpenAI

    client = OpenAI(
        api_key=settings.meeting_asr_api_key,
        base_url=settings.meeting_asr_base_url,
        timeout=max(settings.llm_timeout, 300),
    )
    with path.open("rb") as audio:
        try:
            response = client.audio.transcriptions.create(model=settings.meeting_asr_model, file=audio)
        except Exception as exc:  # noqa: BLE001
            message = str(exc)
            if "no audio segment found" in message or "'1210'" in message:
                raise ValueError("音频中未检测到人声，无法转写；请上传有效的会议录音") from exc
            if "1214" in message or "不支持当前文件格式" in message:
                raise ValueError("音频格式不受 ASR 服务支持，请上传 mp3/wav/m4a 等标准音频文件") from exc
            raise
    text = getattr(response, "text", "") or ""
    if not text.strip():
        raise ValueError("ASR 未返回转录文本")
    result = parse_transcript(text)
    result["asrModel"] = settings.meeting_asr_model
    return result


def _parse_pdf(path: Path) -> dict:
    from pypdf import PdfReader

    pages = []
    segments = []
    for index, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(text)
            segments.append({"text": text, "location": f"PDF第{index}页", "page": index})
    if not pages:
        raise ValueError("PDF 未提取到文本，可能为扫描件")
    return {"text": "\n\n".join(pages), "segments": segments}


def _parse_docx(path: Path) -> dict:
    from docx import Document

    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if text:
                paragraphs.append(text)
    if not paragraphs:
        raise ValueError("Word 未提取到文本")
    return {
        "text": "\n\n".join(paragraphs),
        "segments": [{"text": text, "location": f"Word第{index}段"} for index, text in enumerate(paragraphs, 1)],
    }


def _parse_pptx(path: Path) -> dict:
    from pptx import Presentation

    segments = []
    for slide_no, slide in enumerate(Presentation(str(path)).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"备注：{notes}")
        except (AttributeError, ValueError):
            pass
        if parts:
            segments.append({"text": "\n".join(parts), "location": f"PPT第{slide_no}页", "page": slide_no})
    if not segments:
        raise ValueError("PPTX 未提取到文本或备注")
    return {"text": "\n\n".join(item["text"] for item in segments), "segments": segments}


def _paragraph_result(text: str, label: str) -> dict:
    values = _paragraphs(_normalize(text))
    return {
        "text": "\n\n".join(values),
        "segments": [{"text": value, "location": f"{label}{index}"} for index, value in enumerate(values, 1)],
    }


def _paragraphs(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"\n{2,}|(?<=[。！？!?])\s+", text) if part.strip()]


def _normalize(text: str) -> str:
    return re.sub(r"\r\n?", "\n", text or "").strip()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文本编码不支持")
