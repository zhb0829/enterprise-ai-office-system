"""素材解析、分块、语义向量化与引用检索。"""
from __future__ import annotations

import hashlib
import math
import re
from io import BytesIO

from openai import OpenAI
from sqlalchemy.orm import Session

from ..config import settings
from ..models import MaterialChunk, ReferenceMaterial

SUPPORTED_EXTS = {"md", "txt", "docx", "pdf"}
CHUNK_SIZE = 900
CHUNK_OVERLAP = 120
EMBEDDING_DIMS = 64


class UnsupportedFileError(Exception):
    pass


class MaterialParseError(Exception):
    pass


class EmbeddingError(Exception):
    pass


class EmbeddingClient:
    """通过 OpenAI 兼容接口生成语义向量。"""

    def __init__(self):
        self.is_configured = bool(settings.embedding_api_key and settings.embedding_model)
        self._client = None
        if self.is_configured:
            self._client = OpenAI(
                api_key=settings.embedding_api_key,
                base_url=settings.embedding_base_url,
                timeout=settings.embedding_timeout,
            )

    def embed(self, texts: list[str]) -> list[list[float]]:
        if not self.is_configured:
            raise EmbeddingError("未配置 EMBEDDING_API_KEY 或 EMBEDDING_MODEL")
        try:
            request = {"model": settings.embedding_model, "input": texts}
            if settings.embedding_dimensions:
                request["dimensions"] = settings.embedding_dimensions
            response = self._client.embeddings.create(**request)
        except Exception as exc:  # noqa: BLE001
            raise EmbeddingError(f"Embedding 服务调用失败：{exc}") from exc

        vectors = [[] for _ in texts]
        for item in response.data:
            if 0 <= item.index < len(vectors):
                vectors[item.index] = [float(value) for value in item.embedding]
        if any(not vector for vector in vectors):
            raise EmbeddingError("Embedding 服务返回的向量数量不完整")
        dimensions = {len(vector) for vector in vectors}
        if len(dimensions) != 1:
            raise EmbeddingError("Embedding 服务返回的向量维度不一致")
        expected_dimensions = settings.embedding_dimensions
        actual_dimensions = next(iter(dimensions))
        if expected_dimensions and actual_dimensions != expected_dimensions:
            raise EmbeddingError(
                f"Embedding 维度不符合配置：期望 {expected_dimensions}，实际 {actual_dimensions}"
            )
        return vectors


def parse_material(filename: str, data: bytes) -> str:
    """解析上传文件为纯文本。"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if ext in ("md", "txt"):
            return _decode_text(data)
        if ext == "docx":
            return _ensure_text(_parse_docx(data), "DOCX 未提取到可用文本")
        if ext == "pdf":
            return _ensure_text(_parse_pdf(data), "PDF 未提取到可用文本，可能是扫描件或加密文件")
    except MaterialParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise MaterialParseError(f"素材解析失败: {exc}") from exc
    raise UnsupportedFileError(f"不支持的文件格式: .{ext}，仅支持 md/txt/docx/pdf")


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise MaterialParseError("文本编码不支持，请使用 UTF-8 或 GB18030 编码")


def _ensure_text(text: str, message: str) -> str:
    text = text.strip()
    if not text:
        raise MaterialParseError(message)
    return text


def _parse_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    table_rows = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_rows.append(" | ".join(cells))
    return "\n".join([*paragraphs, *table_rows])


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n".join(pages)


def split_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """按段落优先分块，长段落再按固定窗口切分。"""
    text = re.sub(r"\r\n?", "\n", text or "").strip()
    if not text:
        return []

    chunks: list[str] = []
    current = ""
    for paragraph in [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]:
        if len(paragraph) > chunk_size:
            if current:
                chunks.append(current.strip())
                current = ""
            step = max(1, chunk_size - overlap)
            for start in range(0, len(paragraph), step):
                part = paragraph[start : start + chunk_size].strip()
                if part:
                    chunks.append(part)
            continue
        if len(current) + len(paragraph) + 2 <= chunk_size:
            current = f"{current}\n\n{paragraph}".strip()
        else:
            if current:
                chunks.append(current.strip())
            current = paragraph
    if current:
        chunks.append(current.strip())
    return chunks


def _hash_embed(text: str) -> list[float]:
    """离线开发兜底，不作为配置真实 Embedding 后的生产索引。"""
    vector = [0.0] * EMBEDDING_DIMS
    tokens = re.findall(r"[\u4e00-\u9fa5]{2,}|[A-Za-z0-9_.%-]+", text.lower())
    for token in tokens:
        digest = hashlib.sha256(token.encode("utf-8")).digest()
        idx = int.from_bytes(digest[:2], "big") % EMBEDDING_DIMS
        sign = 1.0 if digest[2] % 2 == 0 else -1.0
        vector[idx] += sign
    norm = math.sqrt(sum(v * v for v in vector)) or 1.0
    return [round(v / norm, 6) for v in vector]


def embedding_info(vector: list[float] | None = None) -> dict:
    if settings.embedding_api_key and settings.embedding_model:
        return {
            "embedding_backend": "remote",
            "embedding_model": settings.embedding_model,
            "embedding_dimensions": len(vector or []),
        }
    return {
        "embedding_backend": "hash-fallback",
        "embedding_model": f"hash-{EMBEDDING_DIMS}",
        "embedding_dimensions": EMBEDDING_DIMS,
    }


def embed_texts(texts: list[str], client: EmbeddingClient | None = None) -> list[list[float]]:
    """批量生成向量；配置真实服务后，远端异常会显式失败而非混入哈希向量。"""
    values = [(text or "").strip() for text in texts]
    if not values:
        return []
    if any(not value for value in values):
        raise EmbeddingError("待向量化文本不能为空")

    client = client or EmbeddingClient()
    if not client.is_configured:
        return [_hash_embed(value) for value in values]

    batch_size = max(1, min(settings.embedding_batch_size, 2048))
    vectors: list[list[float]] = []
    for start in range(0, len(values), batch_size):
        vectors.extend(client.embed(values[start : start + batch_size]))
    return vectors


def embed_text(text: str) -> list[float]:
    """单文本兼容入口。新代码优先在入库时调用 embed_texts 批量请求。"""
    return embed_texts([text])[0]


def _embedding_meta(existing: dict | None, vector: list[float]) -> dict:
    return {**(existing or {}), **embedding_info(vector)}


def embedding_is_compatible(meta: dict | None, query_vector: list[float]) -> bool:
    meta = meta or {}
    backend = meta.get("embedding_backend")
    dimensions = int(meta.get("embedding_dimensions") or 0)
    if settings.embedding_api_key and settings.embedding_model:
        return (
            backend == "remote"
            and meta.get("embedding_model") == settings.embedding_model
            and dimensions == len(query_vector)
        )
    return (
        (backend == "hash-fallback" or not backend)
        and (not dimensions or dimensions == len(query_vector))
    )


def build_material_chunks(db: Session, material: ReferenceMaterial) -> int:
    """为素材生成分块与向量，返回分块数量。"""
    db.query(MaterialChunk).filter(MaterialChunk.material_id == material.id).delete()
    chunks = split_text(material.text_content)
    vectors = embed_texts(chunks)
    for index, (text, vector) in enumerate(zip(chunks, vectors)):
        db.add(
            MaterialChunk(
                material_id=material.id,
                chunk_index=index,
                text=text,
                embedding=vector,
                meta=_embedding_meta({"filename": material.filename, "text_length": len(text)}, vector),
            )
        )
    material.status = "已入库"
    return len(chunks)


def reindex_material_embeddings(db: Session, material_id: int | None = None) -> dict:
    query = db.query(MaterialChunk).join(ReferenceMaterial)
    if material_id is not None:
        query = query.filter(MaterialChunk.material_id == material_id)
    chunks = query.order_by(MaterialChunk.id).all()
    vectors = embed_texts([chunk.text for chunk in chunks]) if chunks else []
    material_ids: set[int] = set()
    for chunk, vector in zip(chunks, vectors):
        chunk.embedding = vector
        chunk.meta = _embedding_meta(chunk.meta, vector)
        material_ids.add(chunk.material_id)
    db.commit()
    return {
        "materials": len(material_ids),
        "chunks": len(chunks),
        **embedding_info(vectors[0] if vectors else None),
    }


def search_material_chunks(
    db: Session,
    query: str,
    material_ids: list[int] | None = None,
    limit: int = 6,
) -> list[dict]:
    """引用检索：关键词分数 + 语义向量余弦相似度，返回可引用片段。"""
    query = (query or "").strip()
    if not query:
        return []
    limit = max(1, min(limit, 20))
    q = db.query(MaterialChunk, ReferenceMaterial.filename).join(ReferenceMaterial)
    if material_ids:
        q = q.filter(MaterialChunk.material_id.in_(material_ids))
    rows = q.all()
    if not rows:
        return []

    query_vec = embed_text(query)
    query_terms = set(re.findall(r"[\u4e00-\u9fa5]{2,}|[A-Za-z0-9_.%-]+", query.lower()))
    ranked = []
    for chunk, filename in rows:
        text = chunk.text or ""
        text_lower = text.lower()
        keyword_hits = sum(1 for term in query_terms if term and term in text_lower)
        cosine = (
            _cosine(query_vec, chunk.embedding or [])
            if embedding_is_compatible(chunk.meta, query_vec)
            else 0.0
        )
        score = keyword_hits * 1.5 + cosine
        if score <= 0:
            continue
        ranked.append(
            {
                "id": chunk.id,
                "material_id": chunk.material_id,
                "filename": filename,
                "chunk_index": chunk.chunk_index,
                "text": text,
                "score": round(float(score), 4),
                "meta": chunk.meta or {},
                "created_at": chunk.created_at,
            }
        )
    ranked.sort(key=lambda item: item["score"], reverse=True)
    return ranked[:limit]


def _cosine(left: list[float], right: list[float]) -> float:
    if not left or not right or len(left) != len(right):
        return 0.0
    denom = math.sqrt(sum(v * v for v in left)) * math.sqrt(sum(v * v for v in right))
    if not denom:
        return 0.0
    return sum(a * b for a, b in zip(left, right)) / denom
